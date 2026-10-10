"""
ai_report.py - API for AI results reports.

  GET  /api/ai-report/status   what is configured (never returns the API key)
  POST /api/ai-report          multipart: file=<results CSV>, location, units,
                               format = docx | xlsx | both (default docx),
                               include_data = 1 | 0 (Excel: add the Data tab)
                               -> JSON with one entry per report in "reports"
  GET  /api/ai-report/file/<token>        the report file (Save as / download)
  POST /api/ai-report/open/<token>        open it in Word / Excel (PC version only)
  POST /api/ai-report/copy/<token>        save a copy in the reports folder (PC version only)
  POST /api/ai-report/open-folder         open the reports folder in Explorer (PC version only)

In local mode a copy of each report is also saved to [reports] output_folder (if set).
"""
from __future__ import annotations

import io
import os
import re
import secrets
import time
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from flask import Blueprint, current_app, jsonify, request, send_file

from app.ai_settings import load_ai_settings
from app.services.report_ai import write_report_text
from app.services.report_docx import build_report
from app.services.report_xlsx import XLSX_MIME, build_xlsx_report
from app.services.results_facts import FactsError, compute_facts, prepare

ai_report_bp = Blueprint("ai_report", __name__, url_prefix="/api/ai-report")
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIMES = {".docx": DOCX_MIME, ".xlsx": XLSX_MIME}
KEEP_HOURS = 24  # generated reports are kept this long for Open / Save as
_TOKEN = re.compile(r"^[0-9a-f]{16}$")
NAME_BASE_MAX = 24   # characters of the location in a report file name
PATH_MAX = 250       # Windows limit is 260 for folder + file name; keep a margin


def _report_dir() -> Path:
    from app.config import Config
    d = Path(Config.UPLOAD_DIR) / "reports"
    d.mkdir(parents=True, exist_ok=True)
    now = time.time()
    for f in [*d.glob("*.docx"), *d.glob("*.xlsx"), *d.glob("*.name")]:  # tidy up old ones
        try:
            if now - f.stat().st_mtime > KEEP_HOURS * 3600:
                f.unlink()
        except OSError:
            pass
    return d


def _report_path(token: str):
    if not _TOKEN.match(token or ""):
        return None
    for ext in MIMES:
        p = _report_dir() / f"{token}{ext}"
        if p.exists():
            return p
    return None


def _store(data: bytes, name: str) -> str:
    """Keep the report for Open / Save as under a short name: <16 hex>.docx / .xlsx,
    with the download name in <token>.name (long names would break Windows' 260-character path limit)."""
    token = secrets.token_hex(8)
    d = _report_dir()
    (d / f"{token}{Path(name).suffix.lower()}").write_bytes(data)
    (d / f"{token}.name").write_text(name, encoding="utf-8")
    return token


def _display_name(path: Path) -> str:
    try:
        return (path.with_suffix(".name")).read_text(encoding="utf-8").strip() or path.name
    except OSError:
        return path.name


def fit_name(folder, name: str) -> str:
    """Shorten the file name (keeping the dates and extension) so folder + name stays under PATH_MAX."""
    room = PATH_MAX - len(str(folder)) - 1 if folder else PATH_MAX
    if len(name) <= room:
        return name
    stem, ext = Path(name).stem, Path(name).suffix
    keep = max(room - len(ext), 12)
    tail = stem[-17:] if re.search(r"\d{8}-\d{8}$", stem) else ""   # the _YYYYMMDD-YYYYMMDD part
    head = stem[: max(keep - len(tail) - 1, 1)].rstrip("_")
    return (f"{head}_{tail.lstrip('_')}" if tail else stem[:keep]) + ext


def _mode() -> str:
    return current_app.config.get("APP_MODE", "local")


def read_reference_text(path: Path | None) -> str:
    if not path or not path.exists():
        return ""
    try:
        if path.suffix.lower() == ".docx":
            from docx import Document
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for t in doc.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def report_filename(location: str, facts: dict) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", location or "Sensor").strip("_")[:NAME_BASE_MAX].strip("_") or "Sensor"
    a = facts["period_start"][:10].replace("-", "")
    b = facts["period_end"][:10].replace("-", "")
    return f"{base}_{a}-{b}.docx"


def _can_open() -> bool:
    return _mode() == "local" and hasattr(os, "startfile")


def _save_copy(folder, name: str, data: bytes) -> str:
    """Write a copy into the reports folder; returns the path or a '(not saved: ...)' note."""
    if not folder:
        return ""
    try:
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / fit_name(folder, name)
        target.write_bytes(data)
        return str(target)
    except OSError as exc:
        return f"(not saved: {exc.strerror or exc})"


@ai_report_bp.get("/status")
def status():
    s = load_ai_settings()
    st = s.public_status(_mode())
    st["mode"] = _mode()
    st["can_open"] = _can_open()
    st["formats"] = ["docx", "xlsx", "both"]
    return jsonify(st)


@ai_report_bp.post("")
def make_report():
    up = request.files.get("file")
    if not up or not up.filename:
        return jsonify({"error": "No results CSV was sent."}), 400
    location = (request.form.get("location") or "").strip()[:120] or "Sensor"
    units = (request.form.get("units") or "").strip()[:12]
    try:
        raw = pd.read_csv(io.BytesIO(up.read()), encoding="utf-8-sig", low_memory=False)
        facts = compute_facts(raw, location=location, units=units, source_name=Path(up.filename).name)
        prepared = prepare(raw)
    except FactsError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"Could not read the results CSV: {exc.__class__.__name__}"}), 400

    fmt = (request.form.get("format") or "docx").strip().lower()
    if fmt not in ("docx", "xlsx", "both"):
        return jsonify({"error": "format must be docx, xlsx or both."}), 400
    include_data = (request.form.get("include_data") or "1").strip() not in ("0", "false", "no")

    s = load_ai_settings()
    ref_text = read_reference_text(s.reference_document)
    ref_name = s.reference_document.name if (s.reference_document and s.reference_document.exists()) else None
    text = write_report_text(facts, ref_text, s.api_key, s.model, s.timeout_seconds)
    base = report_filename(location, facts)[:-len(".docx")]

    files = []
    if fmt in ("docx", "both"):
        files.append((base + ".docx", build_report(prepared, facts, text, ref_name)))
    if fmt in ("xlsx", "both"):
        try:
            files.append((base + ".xlsx", build_xlsx_report(raw, facts, text, location, units,
                                                            Path(up.filename).name, include_data=include_data)))
        except Exception as exc:  # keep the Word report if only the Excel part fails
            current_app.logger.exception("Excel report failed")
            if not files:
                return jsonify({"error": f"Could not build the Excel report: {exc}"}), 500

    reports = []
    for name, data in files:
        saved = _save_copy(s.output_folder, name, data) if _mode() == "local" else ""
        token = _store(data, name)
        reports.append({"token": token, "filename": name, "size": len(data), "saved_to": saved,
                        "kind": Path(name).suffix.lstrip(".")})
    note = text.note or (f"{len(text.removed_sentences)} unverified sentence(s) removed" if text.removed_sentences else "")
    first = reports[0]
    return jsonify({**first, "reports": reports, "source": text.source, "note": note, "can_open": _can_open(),
                    "folder": str(s.output_folder) if (_mode() == "local" and s.output_folder) else ""})


@ai_report_bp.get("/file/<token>")
def report_file(token):
    path = _report_path(token)
    if not path:
        return jsonify({"error": "This report is no longer available. Generate it again."}), 404
    return send_file(str(path), mimetype=MIMES[path.suffix.lower()], as_attachment=True, download_name=_display_name(path))


@ai_report_bp.post("/open/<token>")
def report_open(token):
    """Open the report in Word / Excel on this computer (PC / .exe only)."""
    if not _can_open():
        return jsonify({"error": "Open works in the PC version. Use Save as... or Download here."}), 400
    path = _report_path(token)
    if not path:
        return jsonify({"error": "This report is no longer available. Generate it again."}), 404
    try:
        # Word / Excel show the file name, so open a copy with the readable name - from the
        # short Windows temp folder, not from the (long) project folder, to stay under 260 characters.
        import tempfile
        folder = Path(tempfile.gettempdir()) / "ExtremeAnalyticsLab"
        folder.mkdir(parents=True, exist_ok=True)
        nice = folder / fit_name(folder, _display_name(path))
        nice.write_bytes(path.read_bytes())
        os.startfile(str(nice))  # noqa: S606 - opens the default app (Word / Excel)
    except OSError as exc:
        return jsonify({"error": f"Could not open the report: {exc.strerror or exc}"}), 500
    return jsonify({"opened": True})


@ai_report_bp.post("/copy/<token>")
def report_copy(token):
    """Save a copy of the report in the reports folder ([reports] output_folder), PC version only."""
    if _mode() != "local":
        return jsonify({"error": "Saving to the reports folder works in the PC version. Use Save as... here."}), 400
    s = load_ai_settings()
    if not s.output_folder:
        return jsonify({"error": "No reports folder is set - add [reports] output_folder to config.ini."}), 400
    path = _report_path(token)
    if not path:
        return jsonify({"error": "This report is no longer available. Generate it again."}), 404
    saved = _save_copy(s.output_folder, _display_name(path), path.read_bytes())
    if saved.startswith("(not saved"):
        return jsonify({"error": saved.strip("()")}), 500
    return jsonify({"saved_to": saved})


@ai_report_bp.post("/open-folder")
def open_folder():
    """Open the reports folder in Explorer (PC version only)."""
    s = load_ai_settings()
    if not _can_open() or not s.output_folder:
        return jsonify({"error": "Opening the reports folder works in the PC version with [reports] output_folder set."}), 400
    try:
        s.output_folder.mkdir(parents=True, exist_ok=True)
        os.startfile(str(s.output_folder))  # noqa: S606
    except OSError as exc:
        return jsonify({"error": f"Could not open the folder: {exc.strerror or exc}"}), 500
    return jsonify({"opened": str(s.output_folder)})
