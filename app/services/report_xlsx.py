"""
report_xlsx.py - Excel results report (same layout as
info_main/reports_examples/Jensen_Beach_Results_Analysis_report.xlsx).

Tabs: Summary, AI review, Dashboard, Guide, Data gaps, one tab per month,
Events, Daily overview, Monthly metrics, Column descriptor, Data.

All figures are written as values computed here with pandas (not as Excel
formulas), so the report shows its numbers immediately - also in Excel's
Protected View, in file previews and on phones. Charts are PNG images.

    data = build_xlsx_report(raw_results_df, facts, text, location, units, source_name)
"""
from __future__ import annotations

import datetime as dt
import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from app.services import report_xlsx_charts as charts

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# ---------------------------------------------------------------- styles (from the template)
NAVY, ACC, GRAY, YEL, FH, FB, REC, LAB, LINK = ("131357", "382BEE", "F3F3F3", "FFD966", "A4C2F4",
                                                "C9DAF8", "FFF2CC", "D9D9D9", "1155CC")
WR = Alignment(wrap_text=True, vertical="top")
MID = Alignment(vertical="center")
_thin = Side(style="thin", color="D0D0D0")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
BOT = Border(bottom=Side(style="thin", color="BFBFBF"))
NOTE = "5D6975"


def _fill(c):
    return PatternFill("solid", fgColor=c)


def F(sz=10, b=False, color="000000", u=None, i=False):
    return Font(name="Arial", size=sz, bold=b, color=color, underline=u, italic=i)


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and np.isnan(v):
        return None
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime()
    if v is pd.NaT:
        return None
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return None if np.isnan(v) else float(v)
    return v


def put(ws, ref, v, font=None, fl=None, al=None, fmt=None, bd=None, link=None):
    c = ws[ref]
    c.value = _clean(v)
    c.font = font or F()
    if fl:
        c.fill = _fill(fl)
    if al:
        c.alignment = al
    if fmt:
        c.number_format = fmt
    if bd:
        c.border = bd
    if link:
        c.hyperlink = link
        if font is None:
            c.font = F(10, color=LINK, u="single")
    return c


def band(ws, row, c1, c2, color):
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).fill = _fill(color)


def title_bar(ws, text, right, lastcol):
    for r in (1, 2, 3):
        band(ws, r, 2, lastcol, NAVY)
    put(ws, "B2", text, F(18, True, "FFFFFF"), NAVY, MID)
    c = ws.cell(row=2, column=lastcol - 1, value=right)
    c.font = F(10, False, "FFFFFF")
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="right", vertical="center")
    band(ws, 4, 2, lastcol, ACC)
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[4].height = 4


def link_to(sheet, cell="A1"):
    return f"#'{sheet}'!{cell}"


DATA_HDR = 8      # table header row on the data tabs
DATA_ROW0 = DATA_HDR + 1


def data_tab(ws, title, note, lastcol, gen):
    """Same top as every other tab: navy title bar, back link, one-line note."""
    title_bar(ws, title, f"Report generated: {gen}", lastcol)
    put(ws, "B5", "< Back to summary worksheet", link=link_to("Summary"))
    put(ws, "B6", note, F(10, i=True, color=NOTE))
    ws.column_dimensions["A"].width = 3.6


def header_row(ws, row, col0, heads, dark=True, width=None):
    for i, h in enumerate(heads, col0):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F(10, True, "FFFFFF" if dark else "000000")
        c.fill = _fill(NAVY if dark else LAB)
        c.alignment = WR
        c.border = BOX
        if width:
            ws.column_dimensions[L(i)].width = width


def place_png(ws, png: bytes, anchor: str, width_px: int) -> int:
    im = XLImage(io.BytesIO(png))
    r = width_px / im.width
    im.width, im.height = int(im.width * r), int(im.height * r)
    ws.add_image(im, anchor)
    return im.height


# ---------------------------------------------------------------- data preparation
BOOL_COLS = ("ExtremeFlag", "MaterialExtremeFlag", "ConfirmedShiftFlag", "AbnormalChangeFlag", "MaterialChangeFlag")
NUM_COLS = ("Value", "ZScore", "MaterialityScore", "MeanZ_event", "PersistenceCount", "FrequencyCount",
            "LowerLimit", "UpperLimit", "MovingAverage", "RollingStdDev", "ROC", "ROC_ZScore")


def _bool(s: pd.Series) -> pd.Series:
    if s.dtype == bool:
        return s
    return s.astype(str).str.strip().str.lower().isin(["true", "1", "yes"])


def normalize(raw: pd.DataFrame) -> pd.DataFrame:
    """Typed copy of the results file, kept in file order (the app writes it in time order)."""
    d = raw.rename(columns={c: c.strip() for c in raw.columns}).copy()
    d["Timestamp"] = pd.to_datetime(d["Timestamp"], errors="coerce")
    d["Value"] = pd.to_numeric(d["Value"], errors="coerce")
    d = d.dropna(subset=["Timestamp", "Value"]).reset_index(drop=True)
    for c in BOOL_COLS:
        d[c] = _bool(d[c]) if c in d else False
    for c in NUM_COLS:
        d[c] = pd.to_numeric(d[c], errors="coerce") if c in d else np.nan
    for c in ("EventStartTime", "EventEndTime"):
        d[c] = pd.to_datetime(d[c], errors="coerce") if c in d else pd.NaT
    d["ExtremeDirection"] = d["ExtremeDirection"].fillna("none").astype(str) if "ExtremeDirection" in d else "none"
    d["State"] = d["State"].fillna("Normal").astype(str) if "State" in d else "Normal"
    if "Day/Night" not in d or d["Day/Night"].isna().all():
        h = d["Timestamp"].dt.hour
        d["Day/Night"] = np.where((h >= 7) & (h < 19), "Day", "Night")
    d["Day/Night"] = d["Day/Night"].fillna("").astype(str)
    d["MK"] = d["Timestamp"].dt.year * 100 + d["Timestamp"].dt.month
    d["Date"] = d["Timestamp"].dt.normalize()
    return d


STATE_RANK = {"Extreme triggered": 1, "Extreme/elevated": 2, "Persistent": 3, "Confirmed/Materialized Shift": 4}


def action_for(state: str) -> str:
    return "Investigate" if state == "Confirmed/Materialized Shift" else "Review" if state == "Persistent" else "Monitor"


def find_events(d: pd.DataFrame, ivl_min: float) -> pd.DataFrame:
    rows = []
    sub = d[d["EventStartTime"].notna()]
    for _, g in sub.groupby("EventStartTime", sort=True):
        st = max(g["State"], key=lambda s: STATE_RANK.get(s, 0))
        dirs = g["ExtremeDirection"][g["ExtremeDirection"] != "none"]
        dr = dirs.mode().iloc[0] if len(dirs) else "none"
        peak = g["Value"].min() if dr == "lower" else g["Value"].max()
        t0, t1 = g["Timestamp"].min(), g["Timestamp"].max()
        mz = g["MeanZ_event"].dropna()
        rows.append(dict(start=t0, end=t1, hours=(t1 - t0).total_seconds() / 3600 + ivl_min / 60, dir=dr,
                         peak=float(peak), maxz=float(g["ZScore"].abs().max()) if g["ZScore"].notna().any() else None,
                         meanz=float(mz.iloc[-1]) if len(mz) else None, state=st,
                         score=float(g["MaterialityScore"].max()) if g["MaterialityScore"].notna().any() else 0.0,
                         dn=g["Day/Night"].iloc[0], mk=int(g["MK"].iloc[0]), row=int(g.index[0]) + DATA_ROW0))
    cols = ["start", "end", "hours", "dir", "peak", "maxz", "meanz", "state", "score", "dn", "mk", "row"]
    return pd.DataFrame(rows, columns=cols)


def find_gaps(d: pd.DataFrame, ivl_min: float) -> pd.DataFrame:
    t = d["Timestamp"]
    step = t.diff().dt.total_seconds() / 60
    idx = step[step > 3 * ivl_min].index
    rows = [dict(frm=t[i - 1], to=t[i], hours=step[i] / 60, mk=int(d["MK"][i - 1]), dn=d["Day/Night"][i - 1]) for i in idx]
    return pd.DataFrame(rows, columns=["frm", "to", "hours", "mk", "dn"])


def monthly_table(d: pd.DataFrame, ev: pd.DataFrame, gaps: pd.DataFrame, ivl_min: float) -> pd.DataFrame:
    g = d.groupby("MK")
    mt = pd.DataFrame(index=sorted(d["MK"].unique()))
    mt["Readings"] = g.size()
    mt["Days with data"] = g["Date"].nunique()
    mt["Mean value"] = g["Value"].mean()
    mt["Min value"] = g["Value"].min()
    mt["Max value"] = g["Value"].max()
    mt["Extreme readings"] = g["ExtremeFlag"].sum()
    mt["Extreme share"] = mt["Extreme readings"] / mt["Readings"]
    mt["Upper extremes"] = g["ExtremeDirection"].apply(lambda s: int((s == "upper").sum()))
    mt["Lower extremes"] = g["ExtremeDirection"].apply(lambda s: int((s == "lower").sum()))
    mt["Persistent periods"] = g["State"].apply(lambda s: int((s == "Persistent").sum()))
    mt["Confirmed shift periods"] = g["ConfirmedShiftFlag"].sum()
    mt["Sustained deviation (h)"] = (mt["Persistent periods"] + mt["Confirmed shift periods"]) * ivl_min / 60
    mt["Confirmed share"] = mt["Confirmed shift periods"] / mt["Readings"]
    mt["Abnormal changes"] = g["AbnormalChangeFlag"].sum()
    mt["Material changes"] = g["MaterialChangeFlag"].sum()
    mt["Material changes per day"] = mt["Material changes"] / mt["Days with data"]
    mt["Mean materiality score"] = g["MaterialityScore"].mean()
    mt["Max materiality score"] = g["MaterialityScore"].max()
    mt["Max |Z|"] = d.assign(az=d["ZScore"].abs()).groupby("MK")["az"].max()
    mt["Events"] = ev.groupby("mk").size().reindex(mt.index, fill_value=0) if len(ev) else 0
    mt["Data gap hours"] = gaps.groupby("mk")["hours"].sum().reindex(mt.index, fill_value=0.0) if len(gaps) else 0.0
    for col, base in [("Rank: extreme share", "Extreme share"), ("Rank: sustained (h)", "Sustained deviation (h)"),
                      ("Rank: material changes/day", "Material changes per day"), ("Rank: mean materiality", "Mean materiality score")]:
        mt[col] = mt[base].rank(ascending=False, method="min").fillna(0).astype(int)
    for a in ("Investigate", "Review", "Monitor"):
        mt[a] = [int((ev[ev["mk"] == m]["state"].map(action_for) == a).sum()) if len(ev) else 0 for m in mt.index]
    for dr, dn in (("upper", "Day"), ("upper", "Night"), ("lower", "Day"), ("lower", "Night")):
        mt[f"{dr.title()} – {dn}"] = [int(((d["MK"] == m) & (d["ExtremeDirection"] == dr) & (d["Day/Night"] == dn)).sum())
                                     for m in mt.index]
    return mt


def all_row(d: pd.DataFrame, mt: pd.DataFrame, ivl_min: float) -> dict:
    n = len(d)
    days = d["Date"].nunique()
    a = {"Readings": n, "Days with data": days, "Mean value": d["Value"].mean(), "Min value": d["Value"].min(),
         "Max value": d["Value"].max(), "Extreme readings": int(d["ExtremeFlag"].sum())}
    a["Extreme share"] = a["Extreme readings"] / n
    for c in ("Upper extremes", "Lower extremes", "Persistent periods", "Confirmed shift periods", "Abnormal changes",
              "Material changes", "Events", "Data gap hours", "Sustained deviation (h)"):
        a[c] = mt[c].sum()
    a["Confirmed share"] = a["Confirmed shift periods"] / n
    a["Material changes per day"] = a["Material changes"] / max(days, 1)
    a["Mean materiality score"] = d["MaterialityScore"].mean()
    a["Max materiality score"] = d["MaterialityScore"].max()
    a["Max |Z|"] = d["ZScore"].abs().max()
    return a


def dn_stats(s: pd.DataFrame, ivl_min: float) -> dict:
    n = max(len(s), 1)
    return {"Mean value": s["Value"].mean(), "Min value": s["Value"].min(), "Max value": s["Value"].max(),
            "Extreme share": s["ExtremeFlag"].sum() / n, "Upper extremes": int((s["ExtremeDirection"] == "upper").sum()),
            "Lower extremes": int((s["ExtremeDirection"] == "lower").sum()),
            "Sustained deviation (h)": (int((s["State"] == "Persistent").sum()) + int(s["ConfirmedShiftFlag"].sum())) * ivl_min / 60,
            "Confirmed share": s["ConfirmedShiftFlag"].sum() / n, "Mean materiality score": s["MaterialityScore"].mean(),
            "Max |Z|": s["ZScore"].abs().max()}


FMTS = {"Mean value": "0.00", "Min value": "0.00", "Max value": "0.00", "Extreme share": "0.0%", "Confirmed share": "0.0%",
        "Sustained deviation (h)": "0.0", "Material changes per day": "0.0", "Mean materiality score": "0.000",
        "Max materiality score": "0.000", "Max |Z|": "0.00", "Data gap hours": "0.0", "Days with data": "0"}


def recommendations(m: int, d: pd.DataFrame, mt: pd.DataFrame, window: int | None, k_e: float | None) -> list:
    """Rule-based recommendations for one month: (improve, recommendation, why, measure, learn_more, highlight)."""
    s = d[d["MK"] == m]
    row = mt.loc[m]
    R = []
    lo = s[s["ExtremeDirection"] == "lower"]
    up = s[s["ExtremeDirection"] == "upper"]
    lo_n = (lo["Day/Night"] == "Night").mean() if len(lo) else 0
    up_d = (up["Day/Night"] == "Day").mean() if len(up) else 0
    w_txt = f"{window}-reading" if window else "short"
    if row["Confirmed share"] >= 0.15:
        R.append(("Operating level", "● Review what changed when the level shifted",
                  f"Confirmed shift share: {row['Confirmed share']:.0%} of readings "
                  f"({row['Confirmed shift periods'] * IVL[0] / 60:.0f} h). The value stayed away from its baseline for "
                  "N_M readings in a row with a high mean |Z| - check setpoint or schedule changes, sensor drift or an equipment fault.",
                  "Confirmed share", "Confirmed/Materialized Shift", True))
    if up_d >= 0.65 and lo_n >= 0.5 and len(up) and len(lo):
        R.append(("Daily cycle", "● Check the day/night pattern before acting on single extremes",
                  f"Upper extremes happen mostly by day ({up_d:.0%}) and lower extremes mostly at night ({lo_n:.0%}). "
                  f"With a {w_txt} baseline the normal daily swing itself is flagged. If the swing is intended (setback), "
                  "consider Window = 96 (24 h of 15-minute data).", "Extreme share", "Window (W)", False))
    elif row["Extreme share"] >= 0.10:
        R.append(("Sensitivity", "● Consider a longer baseline window",
                  f"Extreme share is {row['Extreme share']:.0%} of readings"
                  + (f", well above the ~5% expected at K_E = {k_e:g}" if k_e else "")
                  + ". A longer Window (96 = 24 h) or a higher K_E would flag fewer routine swings.", "Extreme share", "K_E", False))
    if row["Material changes per day"] >= 3:
        R.append(("Stability", "● Check for hunting or short-cycling",
                  f"Material changes: {row['Material changes per day']:.1f} per day - fast changes that repeat within the "
                  "frequency window. Check valve/damper control loops and compressor or fan staging.",
                  "Material changes per day", "MaterialChangeFlag", True))
    if row["Data gap hours"] >= 6:
        R.append(("Data quality", "● Fix the data gaps",
                  f"Missing data: {row['Data gap hours']:.1f} h in gaps longer than three reading intervals. "
                  "Events inside a gap cannot be detected.", "Data gap hours", "Data gaps", True))
    return R


IVL = [15.0]  # reading interval in minutes (set per report)

COLUMN_DESCRIPTOR = [
    ("Timestamp", "Time of the reading", "—", "Input"), ("Value", "The reading", "—", "Input"),
    ("AbnormalChangeFlag", "Unusually fast change: |ROC_ZScore| ≥ ROC_K", "ROC_K", "Change"),
    ("ConfirmedShiftFlag", "Reading is in a Confirmed/Materialized Shift", "N_M, K_M", "Materiality"),
    ("DeltaLower", "LowerLimit − Value (positive = below the band)", "limits", "Extreme"),
    ("DeltaLower_pct", "DeltaLower as % of LowerLimit", "limits", "Extreme"),
    ("DeltaUpper", "Value − UpperLimit (positive = above the band)", "limits", "Extreme"),
    ("DeltaUpper_pct", "DeltaUpper as % of UpperLimit", "limits", "Extreme"),
    ("DeltaX", "Change from the previous reading", "—", "Change"),
    ("EventEndTime", "End of the event this reading belongs to", "K_P, N_EXIT", "Persistence"),
    ("EventStartTime", "Start of the event this reading belongs to", "K_E", "Persistence"),
    ("ExtremeDirection", "upper / lower / none", "Method, K_E", "Extreme"),
    ("ExtremeFlag", "Reading outside the normal range", "Method, Window, K_E, p, K_IQR", "Extreme"),
    ("FrequencyCount", "Abnormal changes in the last frequency window", "Frequency window", "Change"),
    ("LowerLimit", "Lower edge of the normal range: MA − K_E·σ", "Method, Window, K_E", "Extreme"),
    ("MaterialChangeFlag", "Abnormal change that recurs (FrequencyCount ≥ min)", "Min recurrences", "Change"),
    ("MaterialExtremeFlag", "Extreme that also lasted N_P readings", "N_P", "Materiality"),
    ("MaterialityScore", "0–1 ranking: Wm·magnitude + Wd·duration + Wf·frequency", "Weights, K_E, N_M", "Materiality"),
    ("MeanZ_event", "Average |Z| over the event so far", "—", "Persistence"),
    ("MovingAverage", "Baseline (rolling mean over Window)", "Method, Window, α", "Extreme"),
    ("PersistenceCount", "Readings in the current event", "K_P, N_EXIT", "Persistence"),
    ("ROC", "Rate of change: (x(t) − x(t−h)) / h", "ROC horizon", "Change"),
    ("ROC_MeanBaseline", "Rolling mean of ROC", "ROC baseline window", "Change"),
    ("ROC_StdDev", "Rolling std of ROC", "ROC baseline window", "Change"),
    ("ROC_ZScore", "(ROC − mean) / std", "ROC baseline window", "Change"),
    ("RollingStdDev", "σ used for the limits", "Method, Window", "Extreme"),
    ("State", "Normal, Extreme triggered, Extreme/elevated, Persistent, Confirmed/Materialized Shift, Persistence ends",
     "K_E, K_P, N_P, N_M, K_M, N_EXIT", "Persistence"),
    ("UpperLimit", "Upper edge of the normal range: MA + K_E·σ", "Method, Window, K_E", "Extreme"),
    ("ZScore", "(Value − MA) / σ", "Method, Window", "Extreme"),
    ("Day/Night", "Day or night category of the reading", "time of day", "Input"),
]


# ---------------------------------------------------------------- AI review sheet
def _ai_items(value):
    """Flatten one AI content field into (bold_title, text) lines."""
    out = []
    if value in (None, "", [], {}):
        return out
    if isinstance(value, str):
        return [("", value)]
    if isinstance(value, dict):
        title = value.get("title") or value.get("promise") or ""
        rest = [f"{k.replace('_', ' ').capitalize()}: {v}" if k not in ("text",) else str(v)
                for k, v in value.items() if k not in ("title",) and v not in (None, "", [])]
        if isinstance(value.get("context"), list):
            rest = [str(x) for x in value["context"]] + [str(value.get("hypothesis", ""))] if value.get("hypothesis") else [str(x) for x in value["context"]]
        return [(title, " ".join(r for r in rest if r))]
    if isinstance(value, list):
        for v in value:
            out.extend(_ai_items(v))
    return out


AI_SECTIONS = [("summary_bullets", "Summary"), ("normal_behavior", "Normal behaviour"),
               ("abnormality_notes", "Abnormal periods"), ("gap_note", "Data gaps"),
               ("minor_observations", "Minor observations"), ("changes", "Rate of change"),
               ("volatility", "Volatility"), ("flags_explanation", "Why the engine flagged what it did"),
               ("recommendations_operations", "Recommendations: operations"),
               ("recommendations_analytics", "Recommendations: analytics settings"),
               ("market_review", "Results against the reference document"), ("market_note", ""),
               ("limitations", "Limitations")]


def ai_source_line(text) -> str:
    if text is None:
        return "No AI wording was requested for this report."
    if getattr(text, "source", "") == "openai":
        line = f"Wording by OpenAI ({text.model}); every number was checked against the data."
        if text.removed_sentences:
            line += f" {len(text.removed_sentences)} sentence(s) with unverified numbers were removed."
        return line
    return "Plain built-in wording (AI not used)." + (f" {text.note}" if getattr(text, "note", "") else "")


def write_ai_sheet(wb, text, gen):
    ws = wb.create_sheet("AI review")
    ws.column_dimensions["A"].width = 3.6
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 110
    title_bar(ws, "AI review", f"Report generated: {gen}", 4)
    put(ws, "B5", "< Back to summary worksheet", link=link_to("Summary"))
    put(ws, "B6", ai_source_line(text), F(10, i=True, color=NOTE))
    put(ws, "B7", "Causes are hypotheses to check, not conclusions. The figures on the other tabs are computed by the app.",
        F(9, i=True, color=NOTE))
    r = 9
    content = getattr(text, "content", None) or {}
    for key, title in AI_SECTIONS:
        items = _ai_items(content.get(key))
        if not items:
            continue
        if title:
            band(ws, r, 2, 3, GRAY)
            put(ws, f"B{r}", title, F(11, True), GRAY)
            r += 1
        for t, body in items:
            if t:
                put(ws, f"B{r}", t, F(10, True), al=WR)
            put(ws, f"C{r}", ("● " if not t and key not in ("normal_behavior", "flags_explanation", "gap_note", "market_note") else "") + body,
                F(10), al=WR)
            ws.row_dimensions[r].height = max(15, 15 * (1 + len(body) // 120))
            r += 1
        r += 1
    ws.sheet_view.showGridLines = False
    return ws


# ---------------------------------------------------------------- main entry point
def build_xlsx_report(raw: pd.DataFrame, facts: dict | None, text, location: str, units: str,
                      source_name: str = "", include_data: bool = True) -> bytes:
    d = normalize(raw)
    if len(d) < 2:
        raise ValueError("Need at least 2 readings to build the Excel report.")
    site = location or "Sensor"
    unit = units or ""
    gen = dt.datetime.now().strftime("%m/%d/%Y %H:%M")
    ivl = float(d["Timestamp"].diff().dt.total_seconds().median() / 60) or 15.0
    IVL[0] = ivl
    s_in = (facts or {}).get("settings_in_file") or {}
    window, k_e = s_in.get("window"), s_in.get("k_e")

    ev = find_events(d, ivl)
    gaps = find_gaps(d, ivl)
    mt = monthly_table(d, ev, gaps, ivl)
    allr = all_row(d, mt, ivl)
    months = list(mt.index)
    mname = {m: dt.date(m // 100, m % 100, 1).strftime("%b %Y") for m in months}
    mlong = {m: dt.date(m // 100, m % 100, 1).strftime("%B %Y") for m in months}
    mt["label"] = [mname[m] for m in months]
    recs = {m: recommendations(m, d, mt, window, k_e) for m in months}
    settings_txt = (f"Window {window}" if window else "Window not detected") + (f", K_E {k_e:g}" if k_e else "")

    wb = Workbook()

    # ---------------- Summary
    ws = wb.active
    ws.title = "Summary"
    for col, w in zip("ABCDEFG", [3.6, 27.6, 50.9, 74.75, 11.75, 26.25, 4]):
        ws.column_dimensions[col].width = w
    title_bar(ws, "Results analysis summary: Monthly overview", f"Report generated: {gen}", 6)
    put(ws, "B6", "Hello,", F(11))
    put(ws, "B7", f"Below is a summary of the {len(months)} month{'s' if len(months) != 1 else ''} of readings for {site} "
                  f"({len(d):,} readings, {d['Timestamp'].min():%m/%d/%Y} – {d['Timestamp'].max():%m/%d/%Y}), and our recommendations.", F(11))
    put(ws, "B8", "Use the tabs at the bottom to find more detail about each month, including event-level actions.", F(11))
    put(ws, "B9", "Read the Guide", link=link_to("Guide"))
    put(ws, "C9", "See the charts (go to 'Dashboard')", link=link_to("Dashboard"))
    put(ws, "D9", "Read the AI review (go to 'AI review')", link=link_to("AI review"))
    r = 11
    # AI summary box
    content = getattr(text, "content", None) or {}
    bullets = [b for _, b in _ai_items(content.get("summary_bullets"))][:4]
    if bullets:
        band(ws, r, 1, 6, GRAY)
        ws.merge_cells(f"B{r}:C{r}")
        put(ws, f"B{r}", "AI summary", F(11, True), GRAY, MID)
        put(ws, f"D{r}", ai_source_line(text), F(9, i=True, color=NOTE), GRAY, MID)
        ws.row_dimensions[r].height = 22
        r += 1
        for b in bullets:
            ws.merge_cells(f"B{r}:F{r}")
            put(ws, f"B{r}", "● " + b, F(10), al=WR)
            ws.row_dimensions[r].height = max(16, 15 * (1 + len(b) // 150))
            r += 1
        r += 2
    for m in months:
        band(ws, r, 1, 6, GRAY)
        ws.merge_cells(f"B{r}:C{r}")
        put(ws, f"B{r}", f"{mlong[m]}  /  Sensor: {site}" + (f" ({unit})" if unit else ""), F(11), GRAY, MID)
        put(ws, f"D{r}", f"Rank by mean materiality: {mt.at[m, 'Rank: mean materiality']} out of {len(months)} months  ·  "
                         f"extremes {mt.at[m, 'Extreme share']:.0%} of readings", F(10), GRAY, MID)
        put(ws, f"E{r}", "Status:", F(10), GRAY, MID)
        rl = recs[m]
        put(ws, f"F{r}", f"{len(rl)} recommendation{'s' if len(rl) != 1 else ''}" if rl else "No recommendations for now",
            F(10), YEL if rl else GRAY, MID)
        ws.row_dimensions[r].height = 22
        r += 2
        if rl:
            for c_, t_ in zip("BCDF", ["Improve", "Recommendations", "Why have we recommended this?", "Learn more"]):
                put(ws, f"{c_}{r}", t_, F(10), bd=BOT)
            ws[f"E{r}"].border = BOT
            r += 1
            for imp, rec, why, met, learn, high in rl:
                put(ws, f"B{r}", "        " + imp, F(10), al=WR)
                put(ws, f"C{r}", rec, F(10, True), al=WR)
                put(ws, f"D{r}", why, F(10), al=WR)
                ws.row_dimensions[r].height = 42
                lk = link_to("Data gaps") if learn == "Data gaps" else link_to("Column descriptor")
                put(ws, f"F{r}", learn, link=lk, al=WR)
                r += 1
            r += 1
            put(ws, f"C{r}", f"Show events and benchmarks for this month (go to '{mname[m]}')", link=link_to(mname[m]))
            r += 3
        else:
            put(ws, f"B{r}", "This month is performing well, no opportunities found", F(10))
            r += 1
            put(ws, f"B{r}", f"Show all the benchmarks (go to '{mname[m]}')", link=link_to(mname[m]))
            r += 3
    band(ws, r, 1, 6, GRAY)
    ws.merge_cells(f"B{r}:C{r}")
    put(ws, f"B{r}", "Data quality  /  Gaps in the readings", F(11), GRAY, MID)
    put(ws, f"D{r}", f"{len(gaps)} gaps, {gaps['hours'].sum() if len(gaps) else 0:.1f} hours of missing data", F(10), GRAY, MID)
    put(ws, f"E{r}", "Status:", F(10), GRAY, MID)
    gh = gaps["hours"].sum() if len(gaps) else 0
    put(ws, f"F{r}", "Gaps to review" if gh >= 6 else "OK", F(10), YEL if gh >= 6 else GRAY, MID)
    put(ws, f"B{r + 2}", "Show the gaps (go to 'Data gaps')", link=link_to("Data gaps"))
    ws.sheet_view.showGridLines = False

    # ---------------- AI review
    write_ai_sheet(wb, text, gen)

    # ---------------- Dashboard
    wdb = wb.create_sheet("Dashboard")
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", [3.6] + [11] * 20):
        wdb.column_dimensions[col].width = w
    title_bar(wdb, "Dashboard", f"Report generated: {gen}", 20)
    put(wdb, "B5", "< Back to summary worksheet", link=link_to("Summary"))
    put(wdb, "B6", "Charts for the whole period. The monthly figures behind them are in the table below.", F(10, i=True, color=NOTE))
    T0 = 9
    put(wdb, f"B{T0 - 1}", "Monthly figures", F(11, True))
    th = ["Month", "Extreme share", "Confirmed share", "Sustained deviation (h)", "Material changes per day", "Investigate",
          "Review", "Monitor", "Upper – Day", "Upper – Night", "Lower – Day", "Lower – Night"]
    header_row(wdb, T0, 2, th, dark=False)
    wdb.row_dimensions[T0].height = 28
    for k, m in enumerate(months, T0 + 1):
        for i, h in enumerate(th, 2):
            v = mt.at[m, "label"] if h == "Month" else mt.at[m, h]
            c = wdb.cell(row=k, column=i, value=_clean(v))
            c.font = F()
            c.border = BOX
            c.number_format = FMTS.get(h, "0")
    r = T0 + len(months) + 3
    W = 1150
    for png, title in [(charts.monthly_panels(mt, unit=unit), "By month"),
                       (charts.events_and_directions(mt), "Events and extremes"),
                       (charts.hour_heatmap(d), "Time of day"),
                       (charts.daily_overview(d, gaps, unit=unit), "By day")]:
        put(wdb, f"B{r}", title, F(11, True))
        h = place_png(wdb, png, f"B{r + 1}", W)
        r += int(h / 20) + 4
    wdb.sheet_view.showGridLines = False

    # ---------------- Guide
    wg = wb.create_sheet("Guide")
    title_bar(wg, "Guide", "", 12)
    for col, w in (("A", 3.6), ("B", 28), ("C", 4), ("D", 100)):
        wg.column_dimensions[col].width = w
    put(wg, "B6", f"This report summarises the Extreme & Change analysis of {site} and gives recommendations and benchmarks for each month.")
    put(wg, "B7", "It shows when the readings left their normal range, how long deviations lasted, how often the value changed "
                  "abnormally fast, and how much data is missing.")
    put(wg, "B9", "How to use this report", F(11, True))
    put(wg, "B10", "Use the tabs at the bottom of the screen to move through the workbook.")
    put(wg, "B12", "Tab (at the bottom of the screen)", F(10, True), LAB)
    wg["C12"].fill = _fill(LAB)
    put(wg, "D12", "What is on the tab?", F(10, True), LAB)
    tabs = [("Summary", ["Overview of each month with recommendations, the AI summary, and links to the month tabs."]),
            ("AI review", ["The written analysis: summary, normal behaviour, abnormal periods, changes, recommendations and limitations."]),
            ("Dashboard", ["Charts: monthly measures, events and extremes, time-of-day heatmap and the daily overview."]),
            ("Month tabs", ["● Analysis funnel: extreme share, sustained deviation hours and material changes per day, with the month's rank.",
                            "● Recommendations with this month vs the best and worst month, and more benchmarks (also day vs night).",
                            "● A chart of the readings against the normal range, and every event sorted by materiality."]),
            ("Data gaps", ["Every break in the data longer than three reading intervals, with the hours missing."]),
            ("Events", ["All events in the whole period in one filterable list."]),
            ("Daily overview", ["One row per day: readings, value range, extremes, persistent and confirmed periods, changes."]),
            ("Monthly metrics", ["The monthly figures behind the Summary and month tabs."]),
            ("Column descriptor", ["What every column of the results file means and which parameters drive it."]),
            ("Data", ["The results file as exported by the app." if include_data else "Not included in this report (option off)."])]
    r = 13
    for t_, lines in tabs:
        target = mname[months[0]] if t_ == "Month tabs" else t_
        put(wg, f"B{r}", t_, F(10, True, LINK, "single"), link=link_to(target))
        for ln in lines:
            put(wg, f"D{r}", ln, F(10), al=WR)
            r += 1
        r += 1
    put(wg, f"B{r}", "How are the results calculated?", F(11, True))
    r += 1
    for ln in [f"Each reading is compared with a rolling baseline of the previous readings ({settings_txt}). |Z| ≥ K_E flags an Extreme reading.",
               "An event continues while |Z| ≥ K_P; after N_P readings it is Persistent, after N_M readings with mean |Z| ≥ K_M it is a Confirmed/Materialized Shift.",
               "The rate of change has its own baseline; |Z_ROC| ≥ ROC_K is an Abnormal Change, and repeats within the frequency window make a Material Change.",
               "The Materiality score (0–1) combines size (|Z|), duration and recurrence. Rank 1 = the month with the most issues.",
               "Recommendations are rules applied to these figures; they point to what to check, they are not a diagnosis of the equipment."]:
        put(wg, f"B{r}", ln)
        r += 1
    r += 1
    put(wg, f"B{r}", "Learn more: info_main\\Statistical_calc_parameters\\Statistical calc parameters manual.docx", F(9, i=True, color=NOTE))
    wg.sheet_view.showGridLines = False

    # ---------------- Data gaps
    wgp = wb.create_sheet("Data gaps")
    title_bar(wgp, "Data gaps", f"Report generated: {gen}", 11)
    for i, t_ in enumerate([f"The analysis expects one reading about every {round(ivl)} minutes.",
                            "A gap is any break of more than three reading intervals. No readings, and so no extremes or events, exist inside a gap.",
                            "This tab lists every gap and the time it covers, so missing data is not mistaken for normal operation."], 6):
        put(wgp, f"B{i}", t_)
    put(wgp, "B10", "Summary of gaps", F(11, True))
    put(wgp, "B11", "Number of gaps")
    put(wgp, "D11", len(gaps), F(10, True))
    put(wgp, "B12", "Missing time (hours)")
    put(wgp, "D12", float(gh), F(10, True), fmt="0.0")
    header_row(wgp, 14, 2, ["From (last reading)", "To (next reading)", "Hours missing", "Readings missing (approx.)",
                            "Day/Night at start", "Severity", "Action"], dark=False, width=20)
    wgp.column_dimensions["A"].width = 3.6
    wgp.column_dimensions["H"].width = 40
    for k, g_ in enumerate(gaps.itertuples(index=False), 15):
        sev = "High" if g_.hours >= 24 else "Medium" if g_.hours >= 6 else "Low"
        vals = [g_.frm, g_.to, g_.hours, max(int(round(g_.hours * 60 / ivl)) - 1, 0), g_.dn, sev,
                "Check the logger / network for this period" if sev != "Low" else "Monitor"]
        for i, v in enumerate(vals, 2):
            c = wgp.cell(row=k, column=i, value=_clean(v))
            c.font = F()
            c.border = BOX
        wgp.cell(row=k, column=2).number_format = wgp.cell(row=k, column=3).number_format = "yyyy-mm-dd hh:mm"
        wgp.cell(row=k, column=4).number_format = "0.0"
        if sev != "Low":
            wgp.cell(row=k, column=7).fill = _fill("F4CCCC" if sev == "High" else REC)
    if not len(gaps):
        put(wgp, "B15", "No gaps found - the readings are continuous.", F(10))
    wgp.sheet_view.showGridLines = False

    # ---------------- Month tabs
    BH = ["Mean value", "Min value", "Max value", "Extreme share", "Upper extremes", "Lower extremes", "Sustained deviation (h)",
          "Confirmed share", "Material changes per day", "Mean materiality score", "Max |Z|", "Data gap hours"]
    funnel = [("Extreme readings", "Share of readings", "Extreme share", "Rank: extreme share", "0.0%"),
              ("Sustained deviations", "Hours (Persistent + Confirmed)", "Sustained deviation (h)", "Rank: sustained (h)", "0.0"),
              ("Material changes", "Per day", "Material changes per day", "Rank: material changes/day", "0.0")]
    for m in months:
        wm = wb.create_sheet(mname[m])
        for col, w in zip("ABCDEFGHIJKLMNO", [3.6, 18, 3, 22, 3, 20, 3, 40, 13, 13, 13, 13, 13, 13, 13]):
            wm.column_dimensions[col].width = w
        for rr in (1, 2, 3):
            band(wm, rr, 2, 15, NAVY)
        put(wm, "B2", "We compared each reading with its own rolling baseline using the settings below", F(12, False, "FFFFFF"), NAVY, MID)
        put(wm, "K2", f"Rank by mean materiality: {mt.at[m, 'Rank: mean materiality']} of {len(months)} months (1 = most issues)",
            F(10, False, "FFFFFF"), NAVY, MID)
        put(wm, "B3", f"Sensor: {site}" + (f" ({unit})" if unit else "") + f",  Period: {mlong[m]},  Value range: "
                      f"{mt.at[m, 'Min value']:.1f} – {mt.at[m, 'Max value']:.1f} {unit}", F(10, False, "FFFFFF"), NAVY, MID)
        put(wm, "K3", f"Report generated: {gen}", F(10, False, "FFFFFF"), NAVY, MID)
        band(wm, 4, 2, 15, ACC)
        wm.row_dimensions[4].height = 4
        put(wm, "B5", "< Back to summary worksheet", link=link_to("Summary"))
        put(wm, "B6", "Analysis funnel", F(11, True))
        band(wm, 7, 2, 6, LAB)
        put(wm, "B7", f"Readings in this month: {int(mt.at[m, 'Readings']):,}", F(10), LAB)
        put(wm, "F7", f"vs other months ({len(months)})", F(10), LAB)
        band(wm, 7, 8, 15, GRAY)
        put(wm, "H7", "Recommendations to improve", F(10), GRAY)
        put(wm, "L7", "Why have we recommended this?", F(10), GRAY)
        for k, (t_, lab, met, rk, fm) in enumerate(funnel):
            rr = 9 + 4 * k
            for c in "DEF":
                wm[f"{c}{rr}"].fill = _fill(FH)
                wm[f"{c}{rr + 1}"].fill = _fill(FB)
                wm[f"{c}{rr + 2}"].fill = _fill(FB)
            put(wm, f"D{rr}", "   " + t_, F(10, True), FH)
            put(wm, f"D{rr + 1}", lab, F(9), FB, WR)
            put(wm, f"F{rr + 1}", "Rank vs other months", F(9), FB, WR)
            put(wm, f"D{rr + 2}", mt.at[m, met], F(14, True), FB, fmt=fm)
            put(wm, f"F{rr + 2}", int(mt.at[m, rk]), F(14, True), FB)
            wm.row_dimensions[rr + 1].height = 26
            wm.row_dimensions[rr + 2].height = 24
        put(wm, "I9", "This month", F(9))
        put(wm, "J9", "Best month", F(9))
        put(wm, "K9", "Worst month", F(9))
        if not recs[m]:
            put(wm, "H10", "This month meets the benchmarks - no recommendations", F(10), REC)
        for k, (imp, rec, why, met, learn, high) in enumerate(recs[m]):
            rr = 10 + 3 * k
            wm.merge_cells(f"H{rr}:H{rr + 1}")
            put(wm, f"H{rr}", rec, F(10, True), REC, WR)
            wm[f"H{rr + 1}"].fill = _fill(REC)
            fm = FMTS.get(met, "0.0")
            put(wm, f"I{rr}", mt.at[m, met], F(10, True, "FF0000" if high else "000000"), fmt=fm)
            put(wm, f"J{rr}", mt[met].min(), F(10), fmt=fm)
            put(wm, f"K{rr}", mt[met].max(), F(10), fmt=fm)
            wm.merge_cells(f"L{rr}:O{rr + 1}")
            put(wm, f"L{rr}", why, F(9), al=WR)
        put(wm, "B23", "More benchmarks", F(11, True))
        for col, h in zip("HIJKLMN", ["Measure", "This month", "Best month", "Worst month", "Whole period",
                                      "This month – Day", "This month – Night"]):
            put(wm, f"{col}24", h, F(9, True), "FFFFFF", WR, bd=BOT)
        wm.row_dimensions[24].height = 26
        sm = d[d["MK"] == m]
        dn = {x: dn_stats(sm[sm["Day/Night"] == x], ivl) for x in ("Day", "Night")}
        for k, h in enumerate(BH, 25):
            fm = FMTS.get(h, "0")
            put(wm, f"H{k}", h, F(10), bd=BOT)
            put(wm, f"I{k}", mt.at[m, h], F(10, True), fmt=fm, bd=BOT)
            put(wm, f"J{k}", mt[h].min(), F(10), fmt=fm, bd=BOT)
            put(wm, f"K{k}", mt[h].max(), F(10), fmt=fm, bd=BOT)
            put(wm, f"L{k}", allr.get(h), F(10), fmt=fm, bd=BOT)
            for col, x in (("M", "Day"), ("N", "Night")):
                put(wm, f"{col}{k}", dn[x].get(h, "Not relevant"), F(10), fmt=fm, bd=BOT)
        notes = ["* Best = lowest month, Worst = highest month for each measure (for Mean/Min/Max value: simply the lowest and highest month).",
                 f"* Settings detected in the results: {settings_txt}."]
        if window:
            notes.insert(1, f"* With a {window}-reading rolling window |Z| cannot exceed (W−1)/√W = {(window - 1) / window ** 0.5:.2f}, "
                            "so Max |Z| is similar in every month; use Extreme share and the events instead.")
        for k, t_ in enumerate(notes, 38):
            put(wm, f"H{k}", t_, F(9, i=True, color=NOTE))
        put(wm, "B42", "Readings against the normal range", F(11, True))
        put(wm, "B43", "Grey band = normal range (moving average ± K_E·σ). Dots outside it are Extreme readings: orange = upper, blue = lower.",
            F(9, i=True, color=NOTE))
        png = charts.month_detail(sm, gaps, title=f"{mlong[m]}: readings and normal range" + (f" ({unit})" if unit else ""), unit=unit)
        place_png(wm, png, "B44", 1150)
        E0 = 68
        put(wm, f"B{E0}", "Events in this month", F(11, True))
        put(wm, f"B{E0 + 1}", "View all events", link=link_to("Events", f"B{DATA_HDR}"))
        put(wm, f"B{E0 + 2}", "* Sorted by materiality score (highest first). Action: Investigate = Confirmed shift, Review = Persistent, "
                              "Monitor = shorter events." + (" Click the Action to see the readings." if include_data else ""),
            F(9, i=True, color=NOTE))
        evh = [("B", "Start"), ("D", "End"), ("F", "Duration (h)"), ("H", "Highest state"), ("I", "Direction"), ("J", "Peak value"),
               ("K", "Max |Z|"), ("L", "Mean |Z|"), ("M", "Max materiality"), ("N", "Day/Night"), ("O", "Action")]
        for col, h in evh:
            put(wm, f"{col}{E0 + 4}", h, F(10, True), "FFFFFF", WR, bd=BOT)
        em = ev[ev["mk"] == m].sort_values("score", ascending=False) if len(ev) else ev
        for k, e in enumerate(em.itertuples(index=False), E0 + 5):
            put(wm, f"B{k}", e.start, fmt="yyyy-mm-dd hh:mm")
            put(wm, f"D{k}", e.end, fmt="yyyy-mm-dd hh:mm")
            put(wm, f"F{k}", e.hours, fmt="0.0")
            put(wm, f"H{k}", e.state)
            put(wm, f"I{k}", e.dir)
            put(wm, f"J{k}", e.peak, fmt="0.00")
            put(wm, f"K{k}", e.maxz, fmt="0.00")
            put(wm, f"L{k}", e.meanz, fmt="0.00")
            put(wm, f"M{k}", e.score, fmt="0.000")
            put(wm, f"N{k}", e.dn)
            a = action_for(e.state)
            c = put(wm, f"O{k}", a, link=f"#Data!A{e.row}" if include_data else None)
            if a != "Monitor":
                c.fill = _fill("F4CCCC" if a == "Investigate" else REC)
        if len(em):
            wm.conditional_formatting.add(f"M{E0 + 5}:M{E0 + 4 + len(em)}",
                                          DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A4C2F4"))
        else:
            put(wm, f"B{E0 + 5}", "No events in this month.")
        wm.freeze_panes = "A5"
        wm.sheet_view.showGridLines = False

    # ---------------- Events
    we = wb.create_sheet("Events")
    data_tab(we, "All events", "One row per event (from Extreme triggered until Persistence ends), sorted by start time. "
             "Use the filter buttons in the header row.", 14, gen)
    EH = ["Start", "End", "Duration (h)", "Direction", "Peak value", "Max |Z|", "Mean |Z| (event)", "Highest state",
          "Max materiality score", "Day/Night at start", "Month", "Action"] + (["Readings"] if include_data else [])
    header_row(we, DATA_HDR, 2, EH, width=16)
    for k, e in enumerate(ev.sort_values("start").itertuples(index=False) if len(ev) else [], DATA_ROW0):
        vals = [e.start, e.end, e.hours, e.dir, e.peak, e.maxz, e.meanz, e.state, e.score, e.dn, mname[e.mk], action_for(e.state)]
        for i, v in enumerate(vals, 2):
            c = we.cell(row=k, column=i, value=_clean(v))
            c.font = F()
            c.border = BOX
        we.cell(row=k, column=2).number_format = we.cell(row=k, column=3).number_format = "yyyy-mm-dd hh:mm"
        for col, fm in ((4, "0.00"), (6, "0.00"), (7, "0.00"), (8, "0.00"), (10, "0.000")):
            we.cell(row=k, column=col).number_format = fm
        if include_data:
            c = we.cell(row=k, column=14, value="View readings")
            c.hyperlink = f"#Data!A{e.row}"
            c.font = F(10, color=LINK, u="single")
    we.freeze_panes = f"B{DATA_ROW0}"
    if len(ev):
        we.auto_filter.ref = f"B{DATA_HDR}:{L(1 + len(EH))}{DATA_HDR + len(ev)}"
    we.sheet_view.showGridLines = False

    # ---------------- Daily overview
    wy = wb.create_sheet("Daily overview")
    data_tab(wy, "Daily overview", "One row per day with data. Bars: share of extreme readings (red) and peak materiality (blue).", 13, gen)
    DH = ["Date", "Readings", "Mean value", "Min value", "Max value", "Extreme readings", "Extreme share", "Persistent periods",
          "Confirmed shift periods", "Abnormal changes", "Material changes", "Max materiality score"]
    header_row(wy, DATA_HDR, 2, DH, width=13)
    gd = d.groupby("Date")
    daily = pd.DataFrame({"Readings": gd.size(), "Mean value": gd["Value"].mean(), "Min value": gd["Value"].min(),
                          "Max value": gd["Value"].max(), "Extreme readings": gd["ExtremeFlag"].sum(),
                          "Persistent periods": gd["State"].apply(lambda s: int((s == "Persistent").sum())),
                          "Confirmed shift periods": gd["ConfirmedShiftFlag"].sum(), "Abnormal changes": gd["AbnormalChangeFlag"].sum(),
                          "Material changes": gd["MaterialChangeFlag"].sum(), "Max materiality score": gd["MaterialityScore"].max()})
    daily["Extreme share"] = daily["Extreme readings"] / daily["Readings"]
    for k, (day, row) in enumerate(daily.iterrows(), DATA_ROW0):
        vals = [day] + [row[h] for h in DH[1:]]
        for i, v in enumerate(vals, 2):
            c = wy.cell(row=k, column=i, value=_clean(v))
            c.font = F()
            c.border = BOX
        wy.cell(row=k, column=2).number_format = "ddd yyyy-mm-dd"
        for col, fm in ((4, "0.00"), (5, "0.00"), (6, "0.00"), (8, "0.0%"), (13, "0.000")):
            wy.cell(row=k, column=col).number_format = fm
    last = DATA_HDR + len(daily)
    wy.conditional_formatting.add(f"H{DATA_ROW0}:H{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4CCCC"))
    wy.conditional_formatting.add(f"M{DATA_ROW0}:M{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A4C2F4"))
    wy.freeze_panes = f"C{DATA_ROW0}"
    wy.column_dimensions["B"].width = 16
    wy.sheet_view.showGridLines = False

    # ---------------- Monthly metrics
    wmm = wb.create_sheet("Monthly metrics")
    data_tab(wmm, "Monthly metrics", f"Computed from the results by the app. Reading interval: {ivl:.2f} minutes (median spacing of the timestamps). "
                   "Rank 1 = the month with the most issues. Gap hours count in the month where the gap starts.", 27, gen)
    MH = ["Month", "Readings", "Days with data", "Mean value", "Min value", "Max value", "Extreme readings", "Extreme share",
          "Upper extremes", "Lower extremes", "Persistent periods", "Confirmed shift periods", "Sustained deviation (h)",
          "Confirmed share", "Abnormal changes", "Material changes", "Material changes per day", "Mean materiality score",
          "Max materiality score", "Max |Z|", "Events", "Data gap hours", "Rank: extreme share", "Rank: sustained (h)",
          "Rank: material changes/day", "Rank: mean materiality"]
    header_row(wmm, DATA_HDR, 2, MH, width=14)
    wmm.row_dimensions[DATA_HDR].height = 42
    for k, m in enumerate(months, DATA_ROW0):
        for i, h in enumerate(MH, 2):
            v = mt.at[m, "label"] if h == "Month" else mt.at[m, h]
            c = wmm.cell(row=k, column=i, value=_clean(v))
            c.font = F()
            c.border = BOX
            c.number_format = FMTS.get(h, "0")
    k = DATA_ROW0 + len(months)
    for i, h in enumerate(MH, 2):
        v = "All" if h == "Month" else allr.get(h)
        c = wmm.cell(row=k, column=i, value=_clean(v))
        c.font = F(10, True)
        c.border = BOX
        c.number_format = FMTS.get(h, "0")
    wmm.freeze_panes = f"C{DATA_ROW0}"
    wmm.sheet_view.showGridLines = False

    # ---------------- Column descriptor
    wc = wb.create_sheet("Column descriptor")
    data_tab(wc, "Column descriptor", f"Meaning of every column in {source_name or 'the results file'} and the parameters that drive it. "
             "Full detail: Statistical calc parameters manual.", 6, gen)
    header_row(wc, DATA_HDR, 2, ["Column", "Meaning", "Driven by", "Step"])
    for k, row in enumerate(COLUMN_DESCRIPTOR, DATA_ROW0):
        for i, v in enumerate(row, 2):
            c = wc.cell(row=k, column=i, value=v)
            c.font = F()
            c.border = BOX
            c.alignment = WR
    r0 = DATA_ROW0 + len(COLUMN_DESCRIPTOR) + 2
    put(wc, f"B{r0}", "Settings detected in this results file", F(11, True))
    put(wc, f"B{r0 + 1}", settings_txt)
    for i, w in enumerate([3.6, 24, 70, 30, 14], 1):
        wc.column_dimensions[L(i)].width = w
    wc.sheet_view.showGridLines = False

    # ---------------- Data
    if include_data:
        wd = wb.create_sheet("Data")
        cols = [c.strip() for c in raw.columns]
        data_tab(wd, "Data", f"The results file {source_name or ''} as exported by the app: {len(d):,} readings, "
                 f"{len(cols)} columns. Meaning of each column: see Column descriptor.", min(len(cols), 16), gen)
        put(wd, "B5", "< Back to summary worksheet", link=link_to("Summary"))
        header_row(wd, DATA_HDR, 1, cols)
        wd.row_dimensions[DATA_HDR].height = 30
        out = d.reindex(columns=cols)
        date_cols = [i for i, h in enumerate(cols, 1) if h in ("Timestamp", "EventStartTime", "EventEndTime")]
        for r_, row in enumerate(out.itertuples(index=False), DATA_ROW0):
            for j, v in enumerate(row, 1):
                wd.cell(row=r_, column=j, value=_clean(v))
        for j in date_cols:
            for r_ in range(DATA_ROW0, len(out) + DATA_ROW0):
                wd.cell(row=r_, column=j).number_format = "yyyy-mm-dd hh:mm"
        for i in range(1, len(cols) + 1):
            wd.column_dimensions[L(i)].width = 14
        wd.column_dimensions["A"].width = 18
        wd.freeze_panes = f"B{DATA_ROW0}"
        wd.auto_filter.ref = f"A{DATA_HDR}:{L(len(cols))}{DATA_HDR + len(out)}"

    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
