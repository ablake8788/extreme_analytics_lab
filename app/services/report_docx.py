"""
report_docx.py - builds the Word (.docx) results report.

Layout matches the approved report: title block, Summary, source table,
1. normal behaviour (+ charts), 2. Extremes (one table per abnormality in the
Titanium format), 3. Changes, 4. Volatility, 5. why the engine flagged,
6. Recommendations, 7. Market review, Limitations, About this report.
All tables are filled from the computed facts.
"""
from __future__ import annotations

import io
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

GRAPHITE = RGBColor(0x1F, 0x2A, 0x33)
MUTED = RGBColor(0x5D, 0x69, 0x75)
BLUE_HEX, LINE_HEX = "2F6FED", "D5DBE1"


# ------------------------------------------------------------------ low-level helpers
def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), LINE_HEX)
        b.append(e)
    tblPr.append(b)


def _bottom_rule(paragraph, color=BLUE_HEX, size=12):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), str(size)); bot.set(qn("w:space"), "6"); bot.set(qn("w:color"), color)
    bdr.append(bot); pPr.append(bdr)


def _run(p, text, bold=False, size=10.5, color=None, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


class Builder:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin = sec.bottom_margin = Inches(0.875)
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"; st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 15), ("Heading 2", 12)):
            h = self.doc.styles[name]
            h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = GRAPHITE
            h.paragraph_format.space_before = Pt(14 if size > 12 else 10); h.paragraph_format.space_after = Pt(6)
        self.n_ops = 0

    def p(self, text="", bold_lead=None, size=10.5, color=None, italic=False):
        para = self.doc.add_paragraph()
        if bold_lead:
            _run(para, bold_lead, bold=True, size=size)
        if text:
            _run(para, text, size=size, color=color, italic=italic)
        return para

    def h1(self, t): self.doc.add_heading(t, level=1)
    def h2(self, t): self.doc.add_heading(t, level=2)

    def bullet(self, text, bold_lead=None):
        para = self.doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        if bold_lead:
            _run(para, bold_lead, bold=True)
        _run(para, text)

    def numbered(self, title, text, num_id=None):
        para = self.doc.add_paragraph(style="List Number")
        para.paragraph_format.space_after = Pt(4)
        _run(para, title.rstrip(".") + ". ", bold=True)
        _run(para, text)
        if num_id is not None:  # belongs to a restarted list
            numPr = para._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = num_id
        return para

    def new_numbered_list(self) -> int:
        """Returns a numbering id for a second list that starts again at 1."""
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        style_num_id = self.doc.styles["List Number"].element.pPr.numPr.numId.val
        abstract_id = numbering.num_having_numId(style_num_id).abstractNumId.val
        num = numbering.add_num(abstract_id)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId

    def table(self, head, rows, widths, first_col_bold=False):
        n = len(widths)
        t = self.doc.add_table(rows=0, cols=n)
        _borders(t)
        t.autofit = False
        if head:
            cells = t.add_row().cells
            for i, h in enumerate(head):
                cells[i].text = ""
                _run(cells[i].paragraphs[0], str(h), bold=True, size=9.5, color=GRAPHITE)
                _shade(cells[i], "EDF0F3")
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                bold = first_col_bold and i == 0
                _run(cells[i].paragraphs[0], "" if val is None else str(val), bold=bold, size=9.5, color=GRAPHITE)
                if bold:
                    _shade(cells[i], "F6F8F9")
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit"); trPr.append(cant)
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
                for para in row.cells[i].paragraphs:
                    para.paragraph_format.keep_with_next = True
                    para.paragraph_format.space_after = Pt(0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv(self, rows):
        return self.table(None, rows, [1.8, 4.7], first_col_bold=True)

    def image(self, png: bytes, width_in=6.2, caption=None):
        para = self.doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))
        if caption:
            self.p(caption, size=9, color=MUTED, italic=True)

    def footer(self, text):
        para = self.doc.sections[0].footer.paragraphs[0]
        _run(para, text + "    Page ", size=8, color=MUTED)
        r = para.add_run()
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
            el = OxmlElement(tag)
            if attr:
                el.set(qn("w:fldCharType"), attr)
            else:
                el.set(qn("xml:space"), "preserve"); el.text = "PAGE"
            r._r.append(el)
        r.font.size = Pt(8); r.font.color.rgb = MUTED

    def bytes(self) -> bytes:
        buf = io.BytesIO(); self.doc.save(buf); return buf.getvalue()


# ------------------------------------------------------------------ charts
def _chart_daily(f, units) -> bytes:
    hourly = pd.DataFrame(f["daily_cycle"]["hourly"])
    fig, ax = plt.subplots(figsize=(10, 2.8), dpi=170)
    ax.plot(hourly["hour"], hourly["typical"], color="#2f6fed", lw=1.6, marker="o", ms=3, label="Typical (median)")
    ax.set_xticks(range(0, 24, 2)); ax.set_xlabel("Hour of day"); ax.set_ylabel(units or "Value")
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    ax.grid(axis="y", color="#e6eaee", lw=0.6); ax.legend(fontsize=7, frameon=False, loc="best")
    ax.set_title("Typical daily cycle", fontsize=10, loc="left", color="#1f2a33")
    plt.tight_layout(); buf = io.BytesIO(); plt.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def _chart_overview(d: pd.DataFrame, f, units) -> bytes:
    prof = {h["hour"]: h["typical"] for h in f["daily_cycle"]["hourly"]}
    fig, ax = plt.subplots(figsize=(10, 3.6), dpi=170)
    for a in f["abnormalities"]:
        ax.axvspan(pd.Timestamp(a["start"]), pd.Timestamp(a["end"]), color="#fdf1de", zorder=0)
    for g in f["data_gaps"]:
        ax.axvspan(pd.Timestamp(g["start"]), pd.Timestamp(g["end"]), color="#e6eaee", zorder=0)
    seg = (d["Timestamp"].diff().dt.total_seconds() / 60 > max(60, 4 * f["interval_minutes"])).cumsum()
    for k, s in d.groupby(seg):
        ax.plot(s["Timestamp"], s["Value"], color="#2f6fed", lw=0.8, label="Measured" if k == 0 else None)
    ax.plot(d["Timestamp"], d["Timestamp"].dt.hour.map(prof), color="#8a96a3", lw=0.7, ls="--", label="Typical for hour of day")
    from matplotlib.patches import Patch
    handles, labels = ax.get_legend_handles_labels()
    if f["abnormalities"]:
        handles.append(Patch(color="#fdf1de")); labels.append("Abnormality")
    if f["data_gaps"]:
        handles.append(Patch(color="#e6eaee")); labels.append("Data gap")
    ax.legend(handles, labels, loc="lower left", fontsize=7, ncol=4, frameon=False)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %d"))
    ax.set_ylabel(units or "Value")
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    ax.grid(axis="y", color="#e6eaee", lw=0.6)
    ax.set_title(f"{f['location']}: {f['period_start'][:10]} to {f['period_end'][:10]}", fontsize=10, loc="left", color="#1f2a33")
    plt.tight_layout(); buf = io.BytesIO(); plt.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------ report
def build_report(df_prepared: pd.DataFrame, f: dict, text, reference_name: str | None) -> bytes:
    c, u = text.content, f.get("units", "")
    B = Builder()
    loc = f["location"]
    period = f"{f['period_start'][:10]} to {f['period_end'][:10]}"

    B.p("Titanium Analytics", size=11, color=MUTED)
    title = B.doc.add_paragraph()
    _run(title, "Results Analysis and Recommendations", bold=True, size=20, color=GRAPHITE)
    sub = B.p(f"{loc}{f' ({u})' if u else ''}  |  {period}", size=11, color=MUTED)
    _bottom_rule(sub)
    B.footer(f"Titanium Analytics  |  {loc}, {period}")

    # Summary
    B.h1("Summary")
    B.p("This report reviews an exported results file from the Extreme & Change Analytics Lab using three methods: "
        "Extremes, Changes and Volatility. Each real abnormality is reported as location, condition, normal range, "
        "observed value, magnitude, duration and status.")
    for s in c["summary_bullets"]:
        B.bullet(s)
    s_in = f["settings_in_file"]
    B.kv([["Source file", f.get("source_name") or "Results of the current analysis"],
          ["Period", f"{f['period_start']} to {f['period_end']}"],
          ["Readings", f"{f['readings']:,}, about one every {f['interval_minutes']} minutes"
           + (f" ({', '.join(f'{v:,} {k}' for k, v in f['day_night_counts'].items())})" if f["day_night_counts"] else "")],
          ["Analysis settings in the file",
           f"Window {s_in['window']} (about {s_in['window_hours']} h), K_E {s_in['k_e']}" if s_in["window"] else "Not detected"],
          ["Reference", reference_name or "No reference document configured"]])

    # 1 normal behaviour
    B.h1("1. How this sensor normally behaves")
    B.p(c["normal_behavior"])
    B.table(["Hours", f"Typical value ({u})" if u else "Typical value"],
            [[b["hours"], f"{b['typical_min']}-{b['typical_max']}" if b["typical_min"] != b["typical_max"] else str(b["typical_min"])]
             for b in f["daily_cycle"]["blocks"]], [2.2, 4.3])
    B.image(_chart_daily(f, u), 5.8)
    B.image(_chart_overview(df_prepared, f, u), 6.2,
            "Shaded: abnormalities (orange) and data gaps (grey). Dashed line: typical value for each hour of the day.")

    # 2 extremes
    B.h1("2. Extremes: abnormalities found")
    B.p(f"Each reading was compared with the typical value for the same hour of the day. "
        f"Deviations larger than {f['deviation_threshold']}{u} lasting at least an hour are reported.")
    n = 0
    for g in f["data_gaps"]:
        n += 1
        B.h2(f"Abnormality {n}: Data gap")
        B.kv([["Location", loc], ["Condition", "No readings received"],
              ["Normal", f"A reading about every {f['interval_minutes']} minutes"],
              ["Observed", f"Last reading {g['start']}, next reading {g['end']}"],
              ["Magnitude", f"About {g['missing_readings']} readings missing"],
              ["Duration", g["duration"]], ["Status", "Resolved, readings resumed"]])
    if c.get("gap_note") and f["data_gaps"]:
        B.p(c["gap_note"])
    notes = c.get("abnormality_notes") or []
    for i, a in enumerate(f["abnormalities"]):
        n += 1
        note = notes[i] if i < len(notes) and isinstance(notes[i], dict) else {}
        title = note.get("title") or f"Unusually {'low' if a['direction'] == 'below' else 'high'} value"
        B.h2(f"Abnormality {n}: {title}")
        B.kv([["Location", loc],
              ["Condition", f"Unusually {'low' if a['direction'] == 'below' else 'high'} for the time of day"],
              ["Normal range", f"{a['typical_range']}{u}"],
              ["Observed", f"{a['observed_range']}{u}; {a['observed']}{u} at {a['peak_time']}"],
              ["Magnitude", f"Up to {abs(a['peak_deviation'])}{u} {a['direction']} typical"],
              ["Duration", f"{a['duration']} ({a['start']} to {a['end']}, {a['weekday']})"],
              ["Status", a["status"]]])
        for b in note.get("context", []) or []:
            B.bullet(b)
        if note.get("hypothesis"):
            B.p(note["hypothesis"], bold_lead="Hypothesis to investigate (not a conclusion): ")
    if n == 0:
        B.p("No substantial abnormality was found against the typical daily cycle.")
    if c.get("minor_observations"):
        B.h2("Minor observations")
        for m in c["minor_observations"]:
            B.bullet(m)

    # 3 / 4
    B.h1("3. Changes")
    for b in c["changes"]:
        B.bullet(b)
    B.h1("4. Volatility")
    for b in c["volatility"]:
        B.bullet(b)

    # 5
    B.h1("5. Why the engine flagged what it did")
    comp = f.get("window_comparison") or []
    if comp:
        head = [""] + [f"Window {r['window']} (about {r['hours']} h)" + (" - in file" if r["window"] == s_in["window"] else "")
                       for r in comp]
        rows = [["Extreme readings"] + [f"{r['extreme_readings']:,}" for r in comp],
                ["Non-normal readings"] + [f"{r['non_normal_readings']:,}" for r in comp],
                ["Separate events"] + [f"{r['events']:,}" for r in comp]]
        w = 6.5 / (len(comp) + 1)
        B.table(head, rows, [w] * (len(comp) + 1))
    B.p(c["flags_explanation"])

    # 6
    B.h1("6. Recommendations")
    B.h2("Building operations")
    for r in c["recommendations_operations"]:
        B.numbered(r.get("title", ""), r.get("text", ""))
    B.h2("Analytics engine and app")
    list_id = B.new_numbered_list()
    for r in c["recommendations_analytics"]:
        B.numbered(r.get("title", ""), r.get("text", ""), num_id=list_id)

    # 7
    if c.get("market_review"):
        B.h1("7. Market review: results against the reference document")
        B.table(["Promise in the document", "Today", "After the recommended fixes"],
                [[m.get("promise", ""), m.get("today", ""), m.get("after", "")] for m in c["market_review"]],
                [2.1, 2.2, 2.2])
        if c.get("market_note"):
            B.p(c["market_note"])

    B.h1("Limitations")
    for b in c["limitations"]:
        B.bullet(b)

    B.h1("About this report")
    about = [f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} by the Extreme & Change Analytics Lab.",
             "All tables and numbers were computed by the app from the results data."]
    if text.source == "openai":
        about.append(f"Wording written by OpenAI ({text.model}) from those facts and the reference document.")
        if text.removed_sentences:
            about.append(f"{len(text.removed_sentences)} AI sentence(s) were removed because they contained numbers "
                         f"not found in the data.")
    else:
        about.append(text.note or "Plain wording built by the app (AI not used).")
    for a in about:
        B.p(a, size=9, color=MUTED)
    return B.bytes()
