"""Tests for the AI results report: facts, OpenAI wording guardrails,
settings from the two ini files, the Word report and the API endpoint."""
import io
import json
import os
import sys
from unittest.mock import patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import numpy as np
import pandas as pd
import pytest

from app.ai_settings import load_ai_settings
from app.services.extreme_change_engine import ChangeParams, ExtremeParams, MaterialityWeights, run_analysis
from app.services.report_ai import allowed_numbers, verify, write_report_text
from app.services.report_docx import build_report
from app.services.results_facts import compute_facts, prepare

ENV = ["OPENAI_API_KEY", "OPENAI_MODEL", "APP_REFERENCE_DOCUMENT", "APP_REPORTS_FOLDER", "APP_MODE",
       "APP_PASSWORD", "APP_SECRET_KEY", "APP_CONFIG_LOCAL", "APP_CONFIG"]


@pytest.fixture(autouse=True)
def clean_env(monkeypatch, tmp_path):
    for k in ENV:
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("APP_CONFIG_LOCAL", str(tmp_path / "none.ini"))
    monkeypatch.setenv("APP_CONFIG", str(tmp_path / "none_shared.ini"))
    yield


def results_csv() -> str:
    """14 days of 15-minute data with a daily cycle, a 12-hour gap on day 6
    and an overnight cold period on day 9 (like Aug 10-11)."""
    t = pd.date_range("2025-08-01 00:00", periods=14 * 96, freq="15min")
    hour = t.hour + t.minute / 60
    occupied = (hour >= 10) & (hour < 19)
    v = np.where(occupied, 74.9, 80.0)
    ramp_down = (hour >= 8) & (hour < 10)
    v = np.where(ramp_down, 80 - (hour - 8) / 2 * 5.1, v)
    ramp_up = (hour >= 19) & (hour < 22)
    v = np.where(ramp_up, 74.9 + (hour - 19) / 3 * 5.1, v)
    v = v + np.random.default_rng(1).normal(0, 0.15, len(v))
    s = pd.Series(v, index=t)
    s[(t >= "2025-08-09 22:00") & (t < "2025-08-10 07:00")] = 72.6  # cold night
    s = s[~((t >= "2025-08-06 10:00") & (t < "2025-08-06 22:00"))]      # data gap
    s = (s / 0.28).round() * 0.28
    r = run_analysis(list(s.index), list(s.values), ExtremeParams(window=20), ChangeParams(), MaterialityWeights())
    tb = r.table.copy()
    tb["Day/Night"] = np.where((tb["Timestamp"].dt.hour >= 7) & (tb["Timestamp"].dt.hour < 19), "Day", "Night")
    return tb.to_csv(index=False)


@pytest.fixture(scope="module")
def csv_text():
    return results_csv()


def facts_of(csv_text):
    return compute_facts(pd.read_csv(io.StringIO(csv_text)), "Test Building, AHU 1", "°F", "test.csv")


# ------------------------------------------------------------------ facts
def test_facts_find_gap_and_cold_night(csv_text):
    f = facts_of(csv_text)
    assert f["readings"] == 14 * 96 - 48
    assert len(f["data_gaps"]) == 1 and f["data_gaps"][0]["duration"].startswith("12 h")
    top = f["abnormalities"][0]
    assert top["direction"] == "below" and top["start"].startswith("2025-08-09 22")
    assert abs(top["observed"] - 72.52) < 0.3 and top["status"] == "Resolved"
    assert f["settings_in_file"]["window"] == 20 and f["sensor_step"] == 0.28
    comp = {r["window"]: r for r in f["window_comparison"]}
    assert comp[96]["events"] < comp[20]["events"]


# ------------------------------------------------------------------ guardrails
def test_invented_numbers_are_removed(csv_text):
    f = facts_of(csv_text)
    content = {"changes": [f"Fastest drop was {f['changes']['typical_fastest_drop']}°F per reading. It reached 99.87°F once."]}
    cleaned, removed = verify(content, allowed_numbers(f, ""))
    assert "99.87" not in json.dumps(cleaned) and len(removed) == 1
    assert str(f["changes"]["typical_fastest_drop"]) in cleaned["changes"][0]


def test_no_key_uses_builtin_wording(csv_text):
    t = write_report_text(facts_of(csv_text), "", api_key="")
    assert t.source == "built-in" and "api_key" in t.note


def _openai_reply(content: dict):
    class R:
        status_code = 200
        def json(self):
            return {"choices": [{"message": {"content": json.dumps(content)}}]}
    return R()


def test_openai_wording_used_and_checked(csv_text):
    f = facts_of(csv_text)
    a = f["abnormalities"][0]
    reply = {"summary_bullets": [f"The space ran {abs(a['peak_deviation'])}°F below typical overnight.",
                                 "Energy use rose 37.25% that week."],
             "abnormality_notes": [{"title": "Night setback not followed", "context": [], "hypothesis": "A schedule override is possible."}]}
    with patch("app.services.report_ai.requests.post", return_value=_openai_reply(reply)) as post:
        t = write_report_text(f, "Extremes Changes Volatility", api_key="sk-test", model="gpt-x", timeout=5)
    sent = post.call_args.kwargs["json"]["messages"][1]["content"]
    assert "Extremes Changes Volatility" in sent and str(a["observed"]) in sent
    assert post.call_args.kwargs["headers"]["Authorization"] == "Bearer sk-test"
    assert t.source == "openai" and len(t.content["summary_bullets"]) == 1
    assert len(t.removed_sentences) == 1 and "37.25" in t.removed_sentences[0]
    assert t.content["abnormality_notes"][0]["title"] == "Night setback not followed"
    assert t.content["changes"]  # missing key filled from built-in wording


def test_openai_error_falls_back(csv_text):
    class R:
        status_code = 401
    with patch("app.services.report_ai.requests.post", return_value=R()):
        t = write_report_text(facts_of(csv_text), "", api_key="sk-bad", model="m", timeout=5)
    assert t.source == "built-in" and "rejected the API key" in t.note


# ------------------------------------------------------------------ settings
def test_settings_order_shared_local_env(monkeypatch, tmp_path):
    shared = tmp_path / "config.ini"
    shared.write_text("[openai]\nmodel = shared-model\n[reports]\noutput_folder = results\n")
    local = tmp_path / "config.local.ini"
    local.write_text("[openai]\napi_key = sk-local\n[reports]\noutput_folder = " + str(tmp_path / "mine") + "\n")
    monkeypatch.setenv("APP_CONFIG", str(shared))
    monkeypatch.setenv("APP_CONFIG_LOCAL", str(local))
    s = load_ai_settings()
    assert s.api_key == "sk-local" and s.model == "shared-model" and s.output_folder == tmp_path / "mine"
    monkeypatch.setenv("OPENAI_API_KEY", "sk-env")
    assert load_ai_settings().api_key == "sk-env"
    assert "sk-" not in json.dumps(load_ai_settings().public_status("local"))


def test_exe_reads_local_ini_but_never_logs_in(monkeypatch, tmp_path):
    from app.server_settings import load_settings
    local = tmp_path / "config.local.ini"
    local.write_text("[app]\nmode = server\n[openai]\napi_key = sk-exe\n")
    monkeypatch.setenv("APP_CONFIG_LOCAL", str(local))
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    try:
        assert load_settings().mode == "local"
        assert load_ai_settings().api_key == "sk-exe"
    finally:
        monkeypatch.delattr(sys, "frozen", raising=False)


# ------------------------------------------------------------------ word report + endpoint
def test_word_report_builds(csv_text):
    from docx import Document
    f = facts_of(csv_text)
    data = build_report(prepare(pd.read_csv(io.StringIO(csv_text))), f, write_report_text(f, ""), "Titanium_analytics.docx")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Results Analysis and Recommendations" in text
    assert "Abnormality 1: Data gap" in text and "No readings received" in text
    assert str(f["abnormalities"][0]["observed"]) in text
    assert len(doc.inline_shapes) == 2


def test_endpoint_generate_open_and_save(monkeypatch, tmp_path, csv_text):
    from app import create_app
    import os as _os
    out = tmp_path / "results"
    monkeypatch.setenv("APP_REPORTS_FOLDER", str(out))
    client = create_app().test_client()
    st = client.get("/api/ai-report/status").get_json()
    assert st["ai_configured"] is False and st["saves_copy"] is True and st["mode"] == "local"
    # 1. generate: JSON with a token, copy in the reports folder
    r = client.post("/api/ai-report", data={"file": (io.BytesIO(csv_text.encode()), "res.csv"),
                                            "location": "Test Building, AHU 1", "units": "°F"},
                    content_type="multipart/form-data")
    assert r.status_code == 200
    d = r.get_json()
    assert d["filename"].startswith("Test_Building_AHU_1_Results_Analysis_20250801-")
    assert (out / d["filename"]).exists() and d["source"] == "built-in" and len(d["token"]) == 32
    # 2. open in Word (PC): calls os.startfile on the kept file
    opened = []
    monkeypatch.setattr(_os, "startfile", lambda p: opened.append(p), raising=False)
    r = client.post(f"/api/ai-report/open/{d['token']}")
    assert r.status_code == 200 and opened and opened[0].endswith(d["filename"])
    # 3. save as: the browser fetches the file
    r = client.get(f"/api/ai-report/file/{d['token']}")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    # unknown / malformed tokens are refused
    assert client.get("/api/ai-report/file/" + "0" * 32).status_code == 404
    assert client.get("/api/ai-report/file/../../etc").status_code == 404


def test_open_not_available_in_server_mode(monkeypatch, tmp_path, csv_text):
    from app import create_app
    import re as _re
    monkeypatch.setenv("APP_MODE", "server")
    monkeypatch.setenv("APP_PASSWORD", "pw")
    monkeypatch.setenv("APP_SECRET_KEY", "k" * 32)
    c = create_app().test_client()
    page = c.get("/login").data.decode()
    tok = _re.search(r'name="csrf" value="([^"]+)"', page).group(1)
    c.post("/login", data={"password": "pw", "csrf": tok, "next": "/"})
    c.get("/login/check?next=/")
    d = c.post("/api/ai-report", data={"file": (io.BytesIO(csv_text.encode()), "res.csv"), "location": "X"},
               content_type="multipart/form-data").get_json()
    assert d["can_open"] is False
    assert c.post(f"/api/ai-report/open/{d['token']}").status_code == 400
    assert c.get(f"/api/ai-report/file/{d['token']}").data[:2] == b"PK"


def test_endpoint_rejects_bad_csv():
    from app import create_app
    client = create_app().test_client()
    r = client.post("/api/ai-report", data={"file": (io.BytesIO(b"a,b\n1,2\n"), "x.csv")},
                    content_type="multipart/form-data")
    assert r.status_code == 400 and "Timestamp" in r.get_json()["error"]
